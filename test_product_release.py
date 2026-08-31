import os

os.environ.setdefault("CONTENT_TTL_SECONDS", "86400")

from academic_voice import analyse_academic_voice
from correction_plan import build_correction_plan, compare_revision_results, _reference_audit
from privacy_lifecycle import build_report_package, purge_result_content
from source_risk import assess_source_risks
from document_correction_pack import build_annotated_document, build_tracked_changes_document
from reference_formatter import format_reference
from payment_control import default_mode
from evidence_resolution import (
    assess_candidate_context_fit,
    build_claim_fingerprint,
    build_document_topic_profile,
    normalise_verification_status,
)
from certificate_builder import build_citation_integrity_certificate
from docx import Document
import io
import zipfile
from pathlib import Path


def sample_result():
    return {
        "main_text": (
            "It is important to note that digital transformation plays a crucial role in performance. "
            "Research shows that integrated systems significantly increase transparency."
        ),
        "missing_in_references": [{"citation": "Adam, 2024"}],
        "uncited_references": [{"reference": "Boateng (2020). Example."}],
        "summary": {"in_text_citations_found": 1, "reference_entries_found": 1},
        "acii": {"score": 62},
    }


def test_voice_review_is_not_authorship_detector():
    review = analyse_academic_voice(sample_result()["main_text"])
    assert review["authorship_inference"] is False
    assert review["signals"]


def test_correction_plan_prioritises_missing_reference():
    plan = build_correction_plan(sample_result())
    assert plan["counts"]["critical"] >= 1
    assert plan["items"][0]["priority"] == "critical"


def test_purge_removes_content_but_keeps_summary():
    purged = purge_result_content(sample_result())
    assert "main_text" not in purged
    assert purged["summary"]["in_text_citations_found"] == 1
    assert purged["privacy"]["content_deleted"] is True


def test_report_package_is_zip():
    data = build_report_package("job-123", sample_result(), extra_files={
        "CiteIntegrity_Annotated_Manuscript.docx": b"annotated",
        "CiteIntegrity_Track_Changes.docx": b"tracked",
    })
    assert data[:2] == b"PK"
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert set(archive.namelist()) == {
            "CiteIntegrity_Annotated_Manuscript.docx",
            "CiteIntegrity_Track_Changes.docx",
            "correction_plan.csv",
            "submission_readiness_report.html",
        }


def test_claim_support_table_has_direct_correction_controls():
    html = Path("templates/new_results.html").read_text(encoding="utf-8")
    assert "Evidence Resolution Workspace" in html
    assert "claim-find-source" in html
    assert "claim-approve-source" in html
    assert "claim-revise" in html
    assert "Add selected citation" in html


def test_revision_comparison():
    original = sample_result()
    revised = {"summary": original["summary"], "acii": {"score": 80}}
    comparison = compare_revision_results(original, revised)
    assert comparison["change"]["total_resolved"] >= 1
    assert comparison["change"]["acii_change"] == 18


def test_not_found_source_is_qualified_not_labelled_fake():
    review = assess_source_risks({"online_verification":{"rows":[{"status":"not_found","reference":"Example"}]}})
    assert review["risks"][0]["risk"] == "metadata_not_found"
    assert "not mean fabricated" in review["risks"][0]["action"].lower()


def test_decision_removes_item_from_pending_counts():
    result = sample_result()
    result["correction_decisions"] = {"missing_reference-1":{"decision":"resolved"}}
    plan = build_correction_plan(result)
    assert plan["counts"]["decided"] >= 1


def test_annotated_and_track_changes_documents_are_generated():
    doc = Document(); doc.add_paragraph("Adam (2024) reported improved transparency.")
    buf = io.BytesIO(); doc.save(buf)
    plan = {"headline":"Review", "items":[{
        "id":"reference_metadata-1", "priority":"important", "evidence":"Adam (2024)",
        "original_text":"Adam (2024)", "proposed_replacement":"Adam and Boateng (2024)",
        "what_is_wrong":"Incomplete author list", "why_it_matters":"Retrieval", "recommended_action":"Correct it",
        "confidence":.99, "decision":"accepted", "auto_apply_allowed":True,
    }]}
    annotated = build_annotated_document(buf.getvalue(), plan)
    tracked, manifest = build_tracked_changes_document(buf.getvalue(), plan)
    assert annotated[:2] == b"PK" and tracked[:2] == b"PK"
    assert manifest["applied_count"] == 1


def test_maintenance_is_a_valid_environment_mode(monkeypatch):
    monkeypatch.setenv("GLOBAL_ACCESS_MODE", "maintenance")
    assert default_mode() == "maintenance"


def test_maintenance_gate_and_developer_bypass_are_present():
    source = open("main.py", encoding="utf-8").read()
    assert "async def maintenance_gate" in source
    assert "developer_request_is_authorized(request)" in source
    assert 'status_code=503' in source
    assert 'automatic_reopening' in source


def test_http_error_handler_preserves_basic_auth_challenge():
    source = open("main.py", encoding="utf-8").read()
    assert "headers=exc.headers" in source
    assert 'WWW-Authenticate' in source


def test_developer_full_testing_session_and_button_are_present():
    source = open("main.py", encoding="utf-8").read()
    assert "Open Full Access" in source
    assert "Open Full Review Unlocked" in source
    assert "developer_session_access_level(request)" in source
    assert "create_developer_session_token(level)" in source
    assert 'path="/"' in source
    assert 'httponly=True' in source and 'secure=True' in source


def test_developer_review_entitlement_is_request_scoped():
    source = open("main.py", encoding="utf-8").read()
    assert "DEVELOPER_ACCESS_LEVELS" in source
    assert "_apply_developer_testing_access" in source
    assert '"developer_access_level": access_level' in source
    assert 'result["payment_required"] = False' in source


def test_payment_notices_are_hidden_for_developer_and_temporary_open_access():
    source = open("templates/new_results.html", encoding="utf-8").read()
    assert "shouldHideAccessPanel" in source
    assert "developer_unlocked" in source
    assert "global_open_access" in source
    assert 'global_access_control?.mode === "open_access"' in source


def test_verification_refresh_preserves_correction_guidance_and_active_tab():
    source = open("templates/new_results.html", encoding="utf-8").read()
    assert "mergeResultPreservingCorrections" in source
    assert "restoreActiveResultTab" in source
    assert 'activeResultTab = target' in source


def test_voice_claim_heuristics_are_conservative_and_not_critical_corrections():
    voice_source = open("academic_voice.py", encoding="utf-8").read()
    plan_source = open("correction_plan.py", encoding="utf-8").read()
    assert "paragraph_has_citation" in voice_source
    assert "PRESENT_STUDY_MARKER" in voice_source
    assert "NON_PROSE_MARKERS" in voice_source
    assert '"priority": "optional"' in voice_source
    assert "possible_claim_needing_source_review" not in voice_source
    assert '"possible_claim_needing_source_review"' in plan_source
    assert "type_limits" in voice_source
    assert "voice_items_added >= 12" in plan_source


def test_academic_voice_revision_requires_approval_and_becomes_track_change():
    main_source = open("main.py", encoding="utf-8").read()
    ui_source = open("templates/new_results.html", encoding="utf-8").read()
    plan_source = open("correction_plan.py", encoding="utf-8").read()
    assert '"/api/academic-voice/{job_id}/approve"' in main_source
    assert "Approve revision for Track Changes" in ui_source
    assert '"track_operation": "replace"' in plan_source
    assert '"decision": "accepted"' in plan_source


def test_approved_citation_reference_addition_and_deletion_become_track_changes():
    doc = Document()
    doc.add_paragraph("Digital systems improve transparency.")
    doc.add_paragraph("Boateng (2020). Unused source.")
    buf = io.BytesIO(); doc.save(buf)
    plan = {"items":[
        {"id":"citation-needed-1", "decision":"accepted", "track_operation":"insert_after", "original_text":"Digital systems improve transparency.", "proposed_replacement":"(Adam, 2024)"},
        {"id":"uncited-reference-1", "decision":"accepted", "track_operation":"delete", "original_text":"Boateng (2020). Unused source."},
        {"id":"missing-reference-1", "decision":"accepted", "track_operation":"append_reference", "proposed_replacement":"Adam, A. (2024). Digital transparency. https://doi.org/10.1000/example"},
    ]}
    tracked, manifest = build_tracked_changes_document(buf.getvalue(), plan)
    assert manifest["applied_count"] == 3
    with zipfile.ZipFile(io.BytesIO(tracked)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "<w:ins" in xml and "<w:del" in xml


def test_source_approval_and_ai_status_endpoints_are_present():
    source = open("main.py", encoding="utf-8").read()
    assert '"/api/corrections/{job_id}/sources/{item_id}"' in source
    assert '"/api/ai/status"' in source
    assert "Select and open a scholarly source" in source


def test_context_source_search_handles_current_candidate_metadata_and_unverified_references():
    suggester = open("citation_suggester.py", encoding="utf-8").read()
    main_source = open("main.py", encoding="utf-8").read()
    ui_source = open("templates/new_results.html", encoding="utf-8").read()
    assert "def _safe_candidate_fields" in suggester
    assert 'isinstance(fields, dict)' in suggester
    assert '"source_verification"' in main_source
    assert "suggest_for_unverified" in main_source
    assert "_reference_citation_context" in main_source
    assert "Find and verify reference" in ui_source


def test_candidate_links_actions_and_non_prose_false_positive_filters():
    ui_source = open("templates/new_results.html", encoding="utf-8").read()
    plan_source = open("correction_plan.py", encoding="utf-8").read()
    worker_source = open("worker.py", encoding="utf-8").read()
    voice_source = open("academic_voice.py", encoding="utf-8").read()
    suggester = open("citation_suggester.py", encoding="utf-8").read()
    assert "Open candidate source" in ui_source
    assert "correction-action-status" in ui_source
    assert "Action recorded" in ui_source
    assert "Open the candidate source link before approving it" in ui_source
    assert "speculative" in plan_source and "unique reference year" in plan_source
    assert "table_or_result" in worker_source and "equation_like" in worker_source
    assert "numeric_density" in voice_source
    assert "strict_citation_identity" in suggester


def test_style_aware_references_voice_redlines_and_claim_support_actions():
    main_source = open("main.py", encoding="utf-8").read()
    ui_source = open("templates/new_results.html", encoding="utf-8").read()
    doc_source = open("document_correction_pack.py", encoding="utf-8").read()
    plan_source = open("correction_plan.py", encoding="utf-8").read()
    formatter = open("reference_formatter.py", encoding="utf-8").read()
    assert "_reference_style_for_result" in main_source
    assert "_style_aware_candidate_text" in main_source
    assert "_style_aware_candidate_citation" in main_source
    assert "insert_after_and_append_reference" in main_source
    assert "_normalised_raw_span" in doc_source
    assert "insert_after_and_append_reference" in doc_source
    assert "Revise or qualify claim" in ui_source
    assert "Find supporting source" in ui_source
    assert "Suggested academic-voice revision" in ui_source
    assert '"claim_support": ["find_source", "add_supporting_citation", "revise_claim"]' in plan_source
    assert 'return f"{formatted[0]}, & {formatted[1]}"' in formatter


def test_v188_reference_style_audit_and_tracked_replacement():
    result = sample_result()
    result["selected_style"] = "apa7"
    result["online_verification"] = {"rows": [{
        "status": "verified", "reference": "Adam, A. (2024). digital integrity. Journal of Integrity, 2(1), 1-9. https://doi.org/10.1000/example",
        "matched_authors_full": ["Adam, Anokye Mohammed"], "matched_year": "2024",
        "matched_title": "Digital integrity", "matched_container_title": "Journal of Integrity",
        "matched_volume": "2", "matched_issue": "1", "matched_pages": "1-9",
        "doi": "10.1000/example",
    }]}
    plan = build_correction_plan(result)
    item = next(row for row in plan["items"] if row["category"] == "reference_style")
    assert "https://doi.org/10.1000/example" in item["proposed_replacement"]
    doc = Document(); doc.add_paragraph(item["original_text"])
    buf = io.BytesIO(); doc.save(buf)
    item["decision"] = "accepted"
    tracked, manifest = build_tracked_changes_document(buf.getvalue(), {"items": [item]})
    assert manifest["applied_count"] == 1
    with zipfile.ZipFile(io.BytesIO(tracked)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "<w:del" in xml and "<w:ins" in xml and "<w:i" in xml


def test_v188_incomplete_reference_is_flagged_not_silently_formatted():
    result = sample_result()
    result["selected_style"] = "harvard"
    result["online_verification"] = {"rows": [{"status": "needs_review", "reference": "Unknown source 2020"}]}
    plan = build_correction_plan(result)
    item = next(row for row in plan["items"] if row["category"] == "reference_incomplete")
    assert item["proposed_replacement"] == ""
    assert item["auto_apply_allowed"] is False
    assert "missing_fields" in item["supporting_metadata"]


def test_numeric_reference_formatter_uses_numbered_list_order():
    formatted = format_reference({"authors":["Adam, Anokye"], "year":"2024", "title":"Digital integrity", "source":"Journal", "volume":"2", "issue":"1", "pages":"1-9"}, "numeric_square")
    assert "Digital integrity." in formatted and "2024;2(1):1-9" in formatted


def test_reference_identity_gate_rejects_different_publication_and_preserves_original():
    result = sample_result()
    result["selected_style"] = "apa7"
    result["online_verification"] = {"rows": [{
        "status":"verified",
        "reference":"Babbie, E. (2021). The practice of social research (15th ed.). Cengage Learning.",
        "matched_authors":["dooly", "vinagre"], "matched_year":"2021",
        "matched_title":"Research into practice: Virtual exchange in language teaching and learning",
        "matched_journal":"Language Teaching", "matched_volume":"55", "matched_issue":"3", "matched_pages":"392",
        "matched_doi":"10.1017/example",
    }]}
    audit = _reference_audit(result)[0]
    assert "Babbie, E." in audit["formatted"]
    assert "Dooly" not in audit["formatted"]
    assert audit["identity"]["accepted"] is False
    plan = build_correction_plan(result)
    conflict = next(row for row in plan["items"] if row["category"] == "reference_identity_conflict")
    assert conflict["proposed_replacement"] == ""
    assert "different publication" in conflict["title"].lower()


def test_corporate_authors_are_not_marked_missing():
    result = sample_result()
    result["selected_style"] = "apa7"
    result["online_verification"] = {"rows": [{"status":"needs_review", "reference":"World Bank. (2020). Justice sector reform and digitization in developing countries. World Bank."}]}
    plan = build_correction_plan(result)
    assert not any(row["category"] == "reference_incomplete" for row in plan["items"])
    audit = _reference_audit(result)[0]
    assert audit["formatted"].startswith("World Bank. (2020).")


def test_mapping_incomplete_and_weak_claims_receive_direct_corrections():
    result = sample_result()
    result["claim_support"] = [
        {"citation":"Judicial Service, 2019", "claim":"Claim one", "support_status":"mapping_incomplete"},
        {"citation":"World Bank, 2020", "claim":"Claim two", "support_status":"weak"},
    ]
    plan = build_correction_plan(result)
    claim_items = [row for row in plan["items"] if row["category"] == "claim_support"]
    assert len(claim_items) == 2
    assert len({row["id"] for row in claim_items}) == 2
    assert all("find_source" in row["available_actions"] for row in claim_items)
    titles = {row["title"] for row in claim_items}
    assert "Claim-to-source mapping is incomplete" in titles
    assert "Claim has weak or unclear support" in titles


def test_evidence_resolution_workspace_unifies_all_unresolved_evidence_groups():
    result = sample_result()
    result["claim_support"] = [
        {"citation": "Judicial Service, 2019", "claim": "Claim one", "support_status": "mapping_incomplete"},
        {"citation": "World Bank, 2020", "claim": "Claim two", "support_status": "weak"},
        {"citation": "OECD, 2022", "claim": "Claim three", "support_status": "insufficient_evidence"},
    ]
    result["online_verification"] = {"rows": [
        {"status": "verified", "reference": "Verified source"},
        {"status": "not_found", "reference": "Unresolved source"},
        {"status": "lookup_failed", "reference": "Lookup failed source"},
    ]}
    plan = build_correction_plan(result)
    workspace = plan["evidence_resolution_workspace"]
    groups = {row["key"] for row in workspace["groups"]}
    assert {
        "unsupported_claims", "incomplete_mappings", "weak_support",
        "missing_references", "uncited_references", "unresolved_verification",
    }.issubset(groups)
    assert workspace["counts"]["total"] >= 7


def test_verification_statuses_are_fair_and_all_unresolved_results_enter_plan():
    assert normalise_verification_status("not_found") == "not_found"
    assert normalise_verification_status("timeout") == "lookup_failed"
    assert normalise_verification_status("metadata_mismatch") == "verified_with_metadata_differences"
    result = sample_result()
    result["online_verification"] = {"rows": [
        {"status": "verified", "reference": "A"},
        {"status": "valid_not_indexed", "reference": "B"},
        {"status": "metadata_mismatch", "reference": "C"},
        {"status": "not_found", "reference": "D"},
        {"status": "timeout", "reference": "E"},
    ]}
    plan = build_correction_plan(result)
    rows = [row for row in plan["items"] if row["category"] == "source_verification"]
    assert len(rows) == 3
    assert all("does not mean" in row["why_it_matters"] for row in rows)
    verify_source = Path("verify.py").read_text(encoding="utf-8")
    worker_source = Path("worker.py").read_text(encoding="utf-8")
    assert 'if st == "likely":\n        return "verified"' not in verify_source
    assert 'trusted_statuses = {"verified"}' in worker_source


def test_context_fit_withholds_out_of_topic_candidates_and_preserves_disclaimer():
    manuscript = "Judicial reform in Ghana concerns court delay, case management, access to justice and public confidence."
    profile = build_document_topic_profile(manuscript)
    unrelated = assess_candidate_context_fit(
        "Digital case management may reduce court delay.",
        manuscript,
        {"title": "Cancer and depression treatment outcomes", "abstract_excerpt": "Clinical oncology trial."},
        profile,
    )
    related = assess_candidate_context_fit(
        "Digital case management may reduce court delay.",
        manuscript,
        {"title": "Digital case management and court delay in judicial reform", "abstract_excerpt": manuscript, "relevance": 90},
        profile,
    )
    assert unrelated["score"] < 70
    assert related["score"] >= 70
    assert related["support_decision"] == "not_assessed_from_metadata"
    assert "not proof" in related["warning"].lower()


def test_structured_profile_and_claim_fingerprint_use_full_manuscript_context():
    manuscript = """Digital Justice Reform in Ghana

Abstract
This study examines digital case management and court delay among judges and lawyers in Ghana using survey and administrative data.

Keywords: judicial reform; digital case management; Ghana; court delay

Research Objectives
To assess whether digital case management reduces court delay in Ghana.

Research Questions
Does digital case management reduce court delay among judges in Ghana?

1. Introduction
Digital case management may reduce court delay in Ghana (World Bank, 2020).
"""
    profile = build_document_topic_profile(manuscript)
    fingerprint = build_claim_fingerprint(
        "Digital case management may reduce court delay in Ghana (World Bank, 2020).",
        manuscript,
        profile,
        "1. Introduction",
    )
    assert profile["title"] == "Digital Justice Reform in Ghana"
    assert "digital case management" in profile["keywords"]
    assert "ghana" in profile["geography"]
    assert {"judges", "lawyers"}.issubset(set(profile["population_terms"]))
    assert "law and justice" in profile["disciplines"]
    assert fingerprint["claimed_relationship"] == "causal_or_effect"
    assert fingerprint["location"] == ["ghana"]
    assert fingerprint["required_evidence_type"].startswith("causal")
    assert fingerprint["nearby_citations"] == ["(World Bank, 2020)"]
    assert fingerprint["current_section"] == "1. Introduction"
    assert fingerprint["ai_api_required"] is False


def test_academic_voice_is_optional_and_off_by_default():
    result = sample_result()
    result["academic_voice_review"] = analyse_academic_voice(result["main_text"])
    plan = build_correction_plan(result)
    assert not any(row["category"] == "academic_voice" for row in plan["items"])
    result["academic_voice_settings"] = {"enabled": True}
    enabled_plan = build_correction_plan(result)
    assert any(row["category"] == "academic_voice" for row in enabled_plan["items"])
    upload_html = Path("templates/new_analyse.html").read_text(encoding="utf-8")
    assert 'id="academicVoice"' in upload_html
    assert 'id="academicVoice" checked' not in upload_html


def test_missing_and_uncited_tabs_are_permanent_standalone_views():
    html = Path("templates/new_results.html").read_text(encoding="utf-8")
    assert 'class="tab-btn permanent-result-tab" data-tab="missingPane"' in html
    assert 'class="tab-btn permanent-result-tab" data-tab="uncitedPane"' in html
    assert 'id="missingPane"' in html and 'id="uncitedPane"' in html
    assert "Missing Citations remain visible" in html or "Missing Citations and Uncited References remain visible" in html
    for removed_tab in ("recoveryPane", "claimPane", "citationNeededPane", "enrichmentPane", "suggestionsPane"):
        assert f'data-tab="{removed_tab}"' not in html
        assert f'id="{removed_tab}"' not in html


def test_certificate_reports_evidence_resolution_and_optional_voice_state():
    result = sample_result()
    result["academic_voice_settings"] = {"enabled": False}
    certificate = build_citation_integrity_certificate(result, job_id="job-123")
    summary = certificate["summary"]
    assert certificate["certificate_title"] == "CiteIntegrity Submission-Readiness Certificate"
    assert summary["evidence_total"] >= 2
    assert summary["evidence_pending"] >= 2
    assert summary["academic_voice_status"] == "Not enabled (optional)"
    assert certificate["verified_stamp"] == "REVIEW RECORDED"
    assert certificate["clearance_status"] == "Revision Required"


def test_v200_release_identity_and_homepage_feature_notice():
    env = Path(".env.example").read_text(encoding="utf-8")
    home = Path("templates/new_index.html").read_text(encoding="utf-8")
    output = Path("templates/new_results.html").read_text(encoding="utf-8")
    assert "RELEASE_VERSION=2.0.0-commercial" in env
    assert "Evidence Resolution Workspace" in home
    assert "Optional writing signals, off by default" in home
    assert "Evidence Resolution Workspace" in output
    assert "PRODUCTION_RESULTS-commercial-v2.0.0" in output
    assert "Simple pay-as-you-go pricing" in home


def test_stats_reports_new_feature_use_without_manuscript_content():
    main_source = Path("main.py").read_text(encoding="utf-8")
    stats_html = Path("templates/stats.html").read_text(encoding="utf-8")
    assert "FEATURE_METRIC_KEYS" in main_source
    assert "context_aware_search_requests" in main_source
    assert "candidates_withheld_by_context" in main_source
    assert "academic_voice_approved_revisions" in main_source
    assert "tracked_approvals_current" in main_source
    assert "Evidence Resolution Workspace" in stats_html
    assert "Context-aware source discovery" in stats_html
    assert "Academic Voice and Writing Signals" in stats_html
    assert "Feature statistics contain aggregate counters only" in stats_html
    assert "Missing citations" in stats_html
    assert "Uncited references" in stats_html


def test_content_purge_retains_only_count_telemetry_for_new_features():
    result = sample_result()
    result.update({
        "feature_usage": {"context_aware_search_requests": 2},
        "correction_source_candidates": {"item-1": [{"title": "Private candidate"}]},
        "correction_source_search_reports": {"item-1": {"message": "Private search query"}},
        "academic_voice_revisions": {"revision-1": {"original_text": "Private passage"}},
        "evidence_resolution_workspace": {"counts": {"total": 3}},
    })
    purged = purge_result_content(result)
    assert purged["feature_usage"]["context_aware_search_requests"] == 2
    assert "correction_source_candidates" not in purged
    assert "correction_source_search_reports" not in purged
    assert "academic_voice_revisions" not in purged
    assert "evidence_resolution_workspace" not in purged


def test_no_abstract_candidate_can_be_shown_only_for_manual_full_text_review():
    manuscript = (
        "Digital justice reform in Ghana examines digital case management, court delay, "
        "judges and lawyers. The study asks whether digital case management reduces court delay."
    )
    claim = "Digital case management may reduce court delay among judges in Ghana."
    profile = build_document_topic_profile(manuscript)
    fingerprint = build_claim_fingerprint(claim, manuscript, profile, "Results")
    candidate = assess_candidate_context_fit(
        claim,
        manuscript,
        {
            "title": "Digital case management and court delay among judges in Ghana",
            "abstract_excerpt": "",
            "relevance": 85,
        },
        profile,
        fingerprint,
    )
    assert candidate["passes_context_gate"] is False
    assert candidate["display_for_manual_review"] is True
    assert candidate["review_tier"] == "manual_review_candidate"
    assert candidate["support_decision"] == "not_assessed_from_metadata"


def test_source_search_report_is_preserved_on_the_correction_item():
    result = sample_result()
    result["correction_source_search_reports"] = {
        "missing_reference-1": {
            "state": "completed_no_candidates",
            "message": "No safe candidate was found.",
            "manual_search_links": [{"label": "Search Crossref", "url": "https://search.crossref.org/"}],
        }
    }
    item = next(row for row in build_correction_plan(result)["items"] if row["id"] == "missing_reference-1")
    assert item["source_search_report"]["state"] == "completed_no_candidates"
    assert item["source_search_report"]["manual_search_links"]


def test_find_supporting_source_has_inline_feedback_and_uses_clicked_button():
    html = Path("templates/new_results.html").read_text(encoding="utf-8")
    assert "renderSourceSearchReport" in html
    assert "source-search-feedback" in html
    assert "Why records were withheld" in html
    assert "Continue manually" in html
    assert "findCorrectionSources(btn.dataset.id, btn)" in html
    assert "renderEvidenceCandidate(correction, source, sourceIndex)" in html
    assert "Manual full-text review required" in html
    assert "actionScope.querySelector" in html


def test_source_discovery_collects_provider_diagnostics_without_network():
    import importlib.util
    import sys
    import types
    from urllib.parse import quote

    if importlib.util.find_spec("requests") is None:
        requests_stub = types.ModuleType("requests")
        requests_stub.get = lambda *_args, **_kwargs: None
        requests_stub.exceptions = types.SimpleNamespace(Timeout=TimeoutError)
        requests_stub.utils = types.SimpleNamespace(quote=quote)
        sys.modules["requests"] = requests_stub
    if importlib.util.find_spec("rapidfuzz") is None:
        rapidfuzz_stub = types.ModuleType("rapidfuzz")
        rapidfuzz_stub.fuzz = types.SimpleNamespace(ratio=lambda *_args: 0, token_set_ratio=lambda *_args: 0)
        sys.modules["rapidfuzz"] = rapidfuzz_stub
    import citation_suggester as suggester

    original_openalex = suggester._query_openalex
    original_crossref = suggester._query_crossref
    try:
        suggester._query_openalex = lambda _query, _rows: []
        suggester._query_crossref = lambda _query, _rows: []
        diagnostics = {}
        suggestions = suggester.suggest_from_context(
            "digital case management court delay Ghana judges",
            top_k=3,
            diagnostics=diagnostics,
        )
    finally:
        suggester._query_openalex = original_openalex
        suggester._query_crossref = original_crossref
    assert suggestions == []
    assert set(diagnostics["providers_attempted"]) == {"openalex", "crossref"}
    assert diagnostics["query_count"] >= 1
    assert diagnostics["provider_result_counts"] == {"openalex": 0, "crossref": 0}
    assert diagnostics["outcome"] == "providers_returned_no_records_or_records_failed_relevance_screen"
