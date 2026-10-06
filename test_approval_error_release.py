"""Reproduce DOCX checksum failures and approval errors through the real route."""
import asyncio
import base64
import io
import json
import struct
import zipfile
from datetime import datetime
from pathlib import Path

import pytest
from docx import Document
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import JSONResponse

from document_correction_pack import _load_docx_status, _repack_docx_media_checksums, build_tracked_changes_document
from test_reference_review_release import load_decision_route, document_xml, SOURCE
from test_verification_coverage_release import load_functions


CLAIM = "Effective procurement controls improve accountability and support consistent public service delivery."
REFERENCE = "Smith, J. (2020). Another study. Example Press."


def document_bytes():
    document = Document()
    document.add_paragraph(CLAIM)
    # A real drawing relationship forces python-docx to read the damaged media
    # member, reproducing the supplied manuscript rather than an orphan member.
    pixel = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aC5kAAAAASUVORK5CYII=")
    document.add_picture(io.BytesIO(pixel))
    document.add_paragraph("References")
    document.add_paragraph(REFERENCE)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def damage_checksum(data, name):
    raw = bytearray(data)
    pos = 0
    while True:
        pos = raw.find(b"PK\x01\x02", pos)
        if pos < 0:
            raise AssertionError("ZIP member not found")
        name_size, extra_size, comment_size = struct.unpack_from("<HHH", raw, pos + 28)
        member = raw[pos + 46:pos + 46 + name_size].decode()
        if member == name:
            crc = struct.unpack_from("<I", raw, pos + 16)[0]
            struct.pack_into("<I", raw, pos + 16, crc ^ 1)
            return bytes(raw)
        pos += 46 + name_size + extra_size + comment_size


def result_fixture():
    return {"selected_style": "apa7", "main_text": CLAIM, "references_raw": [REFERENCE],
            "citation_needed_claims": [{"claim": CLAIM, "confidence": .58}]}


def approval_payload():
    return {"item_id": "citation_needed-1", "decision": "accepted", "action": "insert_citation",
            "original_text": CLAIM, "proposed_replacement": "(Smith, 2022)",
            "approved_source": dict(SOURCE, context_fit_confirmed=True),
            "secondary_replacement": SOURCE["formatted_reference"]}


def asgi_post(route, payload):
    namespace = load_functions("main.py", {"http_exception_handler"}, {
        "Request": Request, "HTTPException": HTTPException, "JSONResponse": JSONResponse,
        "now": lambda: datetime(2026, 10, 6).isoformat()})
    app = FastAPI()
    route.__annotations__["request"] = Request
    app.post("/api/corrections/{job_id}/decision")(route)
    app.add_exception_handler(HTTPException, namespace["http_exception_handler"])
    raw = json.dumps(payload).encode()
    messages = []
    async def receive():
        return {"type": "http.request", "body": raw, "more_body": False}
    async def send(message):
        messages.append(message)
    scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": "POST",
             "scheme": "https", "path": "/api/corrections/job/decision", "raw_path": b"/api/corrections/job/decision",
             "query_string": b"", "headers": [(b"content-type", b"application/json")],
             "server": ("test", 443), "client": ("test", 1234), "root_path": ""}
    asyncio.run(app(scope, receive, send))
    status = next(message["status"] for message in messages if message["type"] == "http.response.start")
    body = b"".join(message.get("body", b"") for message in messages if message["type"] == "http.response.body")
    return status, json.loads(body)


def test_only_media_checksum_is_recomputed_and_all_part_bytes_are_preserved():
    good = document_bytes()
    bad = damage_checksum(good, "word/media/image1.png")
    assert zipfile.ZipFile(io.BytesIO(bad)).testzip() == "word/media/image1.png"
    recovered = _repack_docx_media_checksums(bad)
    with zipfile.ZipFile(io.BytesIO(good)) as original, zipfile.ZipFile(io.BytesIO(recovered)) as restored:
        assert restored.testzip() is None and original.namelist() == restored.namelist()
        for name in original.namelist():
            assert original.read(name) == restored.read(name)
    document, loaded, report = _load_docx_status(bad)
    assert loaded and report["media_checksum_repacked"] and document.paragraphs[0].text == CLAIM


@pytest.mark.parametrize("member", ["word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"])
def test_checksum_damage_in_core_document_parts_is_not_bypassed(member):
    bad = damage_checksum(document_bytes(), member)
    _, loaded, report = _load_docx_status(bad)
    assert not loaded and not report["media_checksum_repacked"] and "fresh DOCX copy" in report["reason"]


@pytest.mark.parametrize("raw,reason", [(None, "Upload the manuscript again"), (b"%PDF", "not an editable DOCX"),
                                      (b"PK truncated data", "damaged or incomplete")])
def test_unavailable_and_invalid_original_files_have_actionable_reasons(raw, reason):
    _, loaded, report = _load_docx_status(raw)
    assert not loaded and reason in report["reason"]


def test_source_approval_with_bad_media_checksum_tracks_claim_and_reference():
    result = result_fixture()
    raw = damage_checksum(document_bytes(), "word/media/image1.png")
    route, saved = load_decision_route(result, raw)
    status, response = asgi_post(route, approval_payload())
    assert status == 200 and response["ok"] and saved.called
    tracked, report = build_tracked_changes_document(raw, response["correction_plan"])
    assert report["applied_count"] == 1 and report["source_document"]["media_checksum_repacked"]
    assert document_xml(tracked).count("<w:ins ") == 2
    assert "(Smith, 2022)" in document_xml(tracked)
    assert zipfile.ZipFile(io.BytesIO(tracked)).testzip() is None


def test_rejected_approval_exposes_same_reason_under_detail_and_error_without_save():
    route, saved = load_decision_route(result_fixture(), None)
    status, response = asgi_post(route, approval_payload())
    assert status == 422 and not saved.called
    assert response["detail"] == response["error"]
    assert "Upload the manuscript again" in response["detail"]


def test_unmatched_claim_still_blocks_approval_and_does_not_append_reference():
    document = Document()
    document.add_paragraph("Another claim only.")
    document.add_paragraph("References")
    stream = io.BytesIO()
    document.save(stream)
    route, saved = load_decision_route(result_fixture(), stream.getvalue())
    status, response = asgi_post(route, approval_payload())
    assert status == 422 and "not found uniquely" in response["detail"] and not saved.called


def test_http_error_handler_preserves_authentication_headers():
    namespace = load_functions("main.py", {"http_exception_handler"}, {
        "Request": Request, "HTTPException": HTTPException, "JSONResponse": JSONResponse, "now": lambda: "now"})
    response = asyncio.run(namespace["http_exception_handler"](None,
        HTTPException(401, "Sign in", headers={"WWW-Authenticate": "Basic"})))
    assert response.headers["www-authenticate"] == "Basic"
