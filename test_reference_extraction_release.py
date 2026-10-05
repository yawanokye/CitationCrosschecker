"""End-to-end regressions for the reported Harvard bibliography undercount.

The first four entries are transcribed from the supplied screenshot. Additional
entries in the long-list tests are synthetic parser fixtures, not real sources.
"""
import io
import ast
import re
import textwrap
import unittest
from pathlib import Path

from docx import Document

import engine


SCREENSHOT_REFERENCES = [
    "Aguilera, R. V., Aragón-Correa, J. A., Marano, V., & Tashman, P. A. 2021. The corporate governance of environmental sustainability: A review and proposal for more integrated research. Journal of Management, 47(6): 1468–1497.",
    "Amankwaa, E. F. 2013. Livelihoods in risk: Exploring health and environmental implications of e-waste recycling as a livelihood strategy in Ghana. Journal of Modern African Studies, 51(4): 551–575.",
    "Atasu, A., Van Wassenhove, L. N., & Sarvary, M. 2009. Efficient take-back legislation. Production and Operations Management, 18(3): 243–258.",
    "Bansal, P., & Roth, K. 2000. Why firms become eco-friendly: Ecological responsiveness and organization theory. Academy of Management Journal, 43(4): 717–736.",
]


def long_references():
    refs = list(SCREENSHOT_REFERENCES)
    for i in range(114):
        suffix = chr(97 + i // 26) + chr(97 + i % 26)
        year = 2000 + i % 25
        refs.append(f"Researcher{suffix}, J. A., & Colleague, P. {year}. Synthetic parser fixture {i}. Test Journal, 12(4): 123–130.")
    return refs


def manuscript(refs, layout="paragraphs"):
    doc = Document()
    for ref in refs:
        key = engine.parse_reference_author_year(ref).key
        author, year = key.split("|")
        doc.add_paragraph(f"The study draws on prior evidence ({author.capitalize()}, {year}).")
    doc.add_paragraph("This deliberately absent source needs review (Missingauthor, 1999).")
    doc.add_heading("References", 1)
    if layout == "paragraphs":
        for ref in refs:
            doc.add_paragraph(ref)
    elif layout == "breaks":
        p = doc.add_paragraph()
        for ref in refs:
            p.add_run(ref).add_break()
    elif layout == "flattened":
        doc.add_paragraph(" ".join(refs))
    elif layout == "wrapped":
        for ref in refs:
            year = engine.YEAR_RE.search(ref)
            doc.add_paragraph(ref[:year.start()].rstrip())
            tail = ref[year.start():]
            cut = tail.index(". ") + 2
            doc.add_paragraph(tail[:cut])
            # More than 500 physical lines for 118 references.
            for part in textwrap.wrap(tail[cut:], width=22, break_long_words=False, break_on_hyphens=False):
                doc.add_paragraph(part)
    doc.add_heading("Appendix A", 1)
    doc.add_paragraph("The survey was administered in 2021. This is not a reference.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class ReferenceExtractionTests(unittest.TestCase):
    def check_result(self, refs, layout):
        result = engine.run_crosscheck(manuscript(refs, layout), "test.docx", style="apa", verify_online=False)
        self.assertEqual(len(result["references_raw"]), len(refs))
        self.assertEqual(result["summary"]["reference_entries_found"], len(refs))
        self.assertEqual(result["references_raw"], refs)
        self.assertEqual([engine.parse_reference_author_year(r).key for r in result["references_raw"]],
                         [engine.parse_reference_author_year(r).key for r in refs])
        missing = result["missing_in_references"]
        self.assertEqual(len(missing), 1, missing)
        self.assertIn("Missingauthor", str(missing))
        self.assertNotIn("Appendix A", " ".join(result["references_raw"]))
        self.assertNotIn("Test Journal", result["main_text"])
        return result

    def test_screenshot_entries_are_four_separate_references(self):
        self.check_result(SCREENSHOT_REFERENCES, "paragraphs")

    def test_118_entries_survive_each_word_layout(self):
        for layout in ("paragraphs", "breaks", "flattened", "wrapped"):
            with self.subTest(layout=layout):
                self.check_result(long_references(), layout)

    def test_long_recovery_has_no_500_line_limit(self):
        text = "References\n" + "\n".join(
            f"{ref}\n\n\n\n\n" for ref in long_references()
        ) + "\nAppendix A\nSurvey details in 2021."
        recovered = engine.recover_references_for_verification(text)
        self.assertEqual(len(recovered), 118)

    def test_wrapped_coauthors_and_years_in_titles_do_not_split(self):
        lines = [
            "Aguilera, R. V., Aragón-Correa, J. A., Marano, V.,",
            "& Tashman, P. A. 2021. Corporate sustainability.",
            "Evidence from the 2020 survey. Journal of Management, 47(6): 1468–1497.",
            "Bansal, P., & Roth, K. 2000. Eco-friendly firms. Journal, 43(4): 717–736.",
        ]
        refs = engine._merge_reference_lines(lines)
        self.assertEqual(len(refs), 2)
        self.assertIn("Tashman", refs[0])
        self.assertIn("2020 survey", refs[0])

    def test_mixed_parenthesised_and_bare_year_formats(self):
        refs = [SCREENSHOT_REFERENCES[0], SCREENSHOT_REFERENCES[1].replace("2013.", "(2013)."),
                SCREENSHOT_REFERENCES[2], SCREENSHOT_REFERENCES[3].replace("2000.", "(2000).")]
        self.check_result(refs, "paragraphs")

    def test_publication_year_page_and_doi_tails_are_not_boundaries(self):
        for line in ("3749. doi:10.1002/jmv.27789", "2021. Journal, 12:22–33.",
                     "Results from the survey in 2021.", "Journal of Management, 2021, 47:12–24."):
            self.assertFalse(engine._looks_like_new_apa_reference_start(line), line)

    def test_corporate_bare_year_reference(self):
        lines = ["World Health Organization. 2023. A global report. Geneva: WHO.",
                 "Ghana Statistical Service. 2021. Population census. Accra: GSS."]
        self.assertEqual(len(engine._merge_reference_lines(lines)), 2)

    def test_autofix_and_verification_worker_receive_all_118_entries(self):
        refs = long_references()
        result = engine.run_crosscheck_with_autofix(manuscript(refs), "test.docx", style="apa")
        self.assertEqual(result["references_raw"], refs)
        self.assertEqual(result["summary"]["reference_entries_found"], 118)
        # Execute the production pure input helpers without Redis/DB imports.
        tree = ast.parse(Path(__file__).with_name("worker.py").read_text())
        names = {"_normalise_reference_for_verification", "_normalise_reference_list_for_verification",
                 "_normalise_references_for_verification"}
        functions = [node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name in names]
        namespace = {"re": re, "_worker_style_family": lambda style: "author_year",
                     "recover_references_for_verification": engine.recover_references_for_verification}
        exec(compile(ast.Module(body=functions, type_ignores=[]), "worker.py", "exec"), namespace)
        self.assertEqual(namespace["_normalise_references_for_verification"](result), refs)


if __name__ == "__main__":
    unittest.main()
