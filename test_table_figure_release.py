"""Focused regression checks for submission readiness and tracked approvals."""

import io
import unittest
import zipfile

from docx import Document

from correction_plan import _reference_audit, build_correction_plan
from document_correction_pack import build_tracked_changes_document
from entitlements import apply_entitlements_to_result
from table_figure_audit import audit_tables_figures
from privacy_lifecycle import build_report_package


def word_bytes(doc):
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


class TableFigureReleaseTests(unittest.TestCase):
    def test_structural_audit_distinguishes_captions_from_callouts(self):
        doc = Document()
        doc.add_paragraph("As Table 1 shows, the trial improved outcomes. Figure 3 is discussed below.")
        doc.add_paragraph("Table 1. Patient outcomes", style="Caption")
        doc.add_table(rows=2, cols=2)
        doc.add_paragraph("Figure 2. Study flow", style="Caption")
        audit = audit_tables_figures(word_bytes(doc), "study.docx")
        self.assertEqual(audit["coverage"], "docx_structural")
        self.assertEqual(audit["summary"]["tables"], 1)
        self.assertEqual(audit["summary"]["figures"], 1)
        self.assertEqual(audit["summary"]["unreferenced"], 1)
        self.assertEqual(audit["summary"]["missing_targets"], 1)
        self.assertEqual({f["type"] for f in audit["findings"]}, {"unreferenced", "missing_target"})

    def test_plural_lists_and_small_ranges_count_each_target(self):
        doc = Document()
        doc.add_paragraph("Tables 1 and 2 present the results; Figures 1–3 show the flow.")
        for number in (1, 2):
            doc.add_paragraph(f"Table {number}. Outcomes", style="Caption")
        for number in (1, 2, 3):
            doc.add_paragraph(f"Figure {number}. Flow", style="Caption")
        audit = audit_tables_figures(word_bytes(doc), "study.docx")
        self.assertEqual(audit["summary"]["unreferenced"], 0)
        self.assertEqual(audit["summary"]["missing_targets"], 0)

    def test_missing_reference_metadata_and_approved_case_appear_as_distinct_actions(self):
        result = {
            "selected_style": "apa7",
            "main_text": "According to smith (2023), the intervention helps.",
            "references_raw": ["Smith, J. (2023). A study.",
                               "Brown, A. (2020). Complete study. Nursing Journal, 4(1), 12-18."],
            "correction_decisions": {"approved-source": {
                "decision": "accepted", "approved_source": {"authors": ["Smith, Jane"], "year": 2023}
            }},
        }
        plan = build_correction_plan(result)
        categories = {item["category"] for item in plan["items"]}
        self.assertIn("reference_incomplete", categories)
        self.assertIn("citation_case", categories)
        item = next(item for item in plan["items"] if item["category"] == "citation_case")
        self.assertEqual(item["original_text"], "smith (2023)")
        self.assertEqual(item["proposed_replacement"], "Smith (2023)")
        self.assertFalse(item["auto_apply_allowed"])

    def test_approved_completed_reference_survives_plan_rebuild_for_export(self):
        original = "Smith, J. (2023). A study."
        completed = "Smith, J. (2023). A study. Nursing Journal, 4(1), 12-18."
        result = {"selected_style": "apa7", "main_text": "Smith (2023) argued this.",
                  "references_raw": [original], "correction_decisions": {
                      "reference-incomplete-1": {"decision": "accepted", "action": "replace_reference",
                                                 "proposed_replacement": completed,
                                                 "approved_source": {"title": "A study", "url": "https://doi.org/10.1234/example"}}
                  }}
        item = next(i for i in build_correction_plan(result)["items"] if i["category"] == "reference_incomplete")
        self.assertEqual(item["proposed_replacement"], completed)
        doc = Document(); doc.add_paragraph(original)
        _, manifest = build_tracked_changes_document(word_bytes(doc), {"items": [item]})
        self.assertEqual(manifest["applied_count"], 1)

    def test_numeric_style_suggests_outlier_without_rewriting_consistent_initials(self):
        refs = [
            "1. Smith JA. Care outcomes. Nursing Journal. 2023;4:12-18.",
            "2. Jones AB. More care outcomes. Nursing Journal. 2022;5:20-28.",
            "Brown CD. Care intervention. Nursing Journal. 2021;3:2-9.",
        ]
        audit = _reference_audit({"selected_style": "numeric_superscript", "references_raw": refs})
        self.assertFalse(audit[0]["needs_formatting"])
        self.assertFalse(audit[1]["needs_formatting"])
        self.assertTrue(audit[2]["needs_formatting"])
        self.assertTrue(audit[2]["formatted"].startswith("3. Brown CD."))

    def test_harvard_year_without_parentheses_is_parsed_before_completeness_audit(self):
        refs = ["Smith, J. 2023. Outcomes. Nursing Journal, 4(1), 12-18.",
                "Jones, A. 2021. Findings. Nursing Journal, 5(1), 22-29."]
        audit = _reference_audit({"selected_style": "apa", "references_raw": refs})
        self.assertEqual([row["missing"] for row in audit], [[], []])
        self.assertEqual(audit[0]["style"], "harvard")

    def test_accepted_edits_preserve_word_runs_and_use_revision_markup(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("See ").bold = True
        paragraph.add_run("Table 2").italic = True
        paragraph.add_run(" for details and smith (2023).")
        plan = {"items": [
            {"id": "table", "category": "table_figure", "decision": "accepted",
             "original_text": "Table 2", "proposed_replacement": "Table 1", "track_operation": "replace"},
            {"id": "case", "category": "citation_case", "decision": "accepted",
             "original_text": "smith (2023)", "proposed_replacement": "Smith (2023)", "track_operation": "replace_all"},
        ]}
        updated, manifest = build_tracked_changes_document(word_bytes(doc), plan)
        xml = Document(io.BytesIO(updated)).paragraphs[0]._p.xml
        self.assertEqual(manifest["applied_count"], 2)
        self.assertIn("w:del", xml)
        self.assertIn("w:ins", xml)
        self.assertIn("w:b", xml)
        self.assertIn("w:i", xml)
        self.assertIn("Smith (2023)", xml)

    def test_approved_missing_callout_is_inserted_at_exact_body_anchor(self):
        doc = Document()
        paragraph = doc.add_paragraph(); paragraph.add_run("Patient ").bold = True
        paragraph.add_run("outcomes improved over time.")
        plan = {"items": [{"id": "callout", "category": "table_figure", "decision": "accepted",
                           "original_text": "outcomes improved", "proposed_replacement": "Table 1 shows the trend.",
                           "track_operation": "insert_after"}]}
        output, manifest = build_tracked_changes_document(word_bytes(doc), plan)
        self.assertEqual(manifest["applied_count"], 1)
        xml = Document(io.BytesIO(output)).paragraphs[0]._p.xml
        self.assertIn("w:ins", xml)
        self.assertIn("Table 1 shows the trend.", xml)
        self.assertIn("w:b", xml)

    def test_free_preview_keeps_counts_but_samples_findings(self):
        audit = {"summary": {"needs_review": 20}, "coverage": "docx_structural",
                 "captions": [{"text": "Table 1. Private title"}],
                 "findings": [{"id": f"f-{i}", "evidence": f"private-{i}"} for i in range(20)]}
        safe = apply_entitlements_to_result({"summary": {"reference_entries_found": 2},
                                              "table_figure_audit": audit}, paid=False)
        self.assertEqual(safe["table_figure_audit"]["summary"]["needs_review"], 20)
        self.assertLessEqual(len(safe["table_figure_audit"]["findings"]), 10)
        self.assertEqual(safe["table_figure_audit"]["captions"], [])
        self.assertTrue(safe["table_figure_audit"]["preview_read_only"])

    def test_approved_source_changes_citation_and_adds_reference_before_appendix(self):
        doc = Document()
        doc.add_paragraph("Earlier work by smith (2023) informed the design.")
        doc.add_heading("References", level=1)
        doc.add_paragraph("Jones, A. (2021). Prior work.")
        doc.add_heading("Appendix", level=1)
        doc.add_paragraph("Supplementary text")
        item = {"id": "source", "decision": "accepted", "category": "missing_reference",
                "original_text": "smith (2023)", "proposed_replacement": "Smith (2023)",
                "secondary_replacement": "Smith, J. (2023). Verified work. Journal of Care.",
                "track_operation": "replace_citation_and_append_reference"}
        output, status = build_tracked_changes_document(word_bytes(doc), {"items": [item]})
        body = Document(io.BytesIO(output))
        self.assertEqual(status["applied"], ["source"])
        self.assertIn("w:del", body.paragraphs[0]._p.xml)
        self.assertIn("Smith (2023)", body.paragraphs[0]._p.xml)
        xml = body._element.body.xml
        self.assertLess(xml.index("Verified work"), xml.index("Appendix"))
        self.assertIn("w:ins", xml)
        self.assertNotIn("change-control note", xml)

    def test_table_introduction_is_editable_and_tracked_before_caption(self):
        doc = Document()
        doc.add_heading("Findings", level=1)
        doc.add_paragraph("Patient outcomes improved over time.")
        doc.add_paragraph("Table 1. Patient outcomes", style="Caption")
        doc.add_table(rows=1, cols=1)
        findings = audit_tables_figures(word_bytes(doc), "study.docx")["findings"]
        issue = next(f for f in findings if f["type"] == "unreferenced")
        item = {"id": "intro", "decision": "accepted", "category": "table_figure",
                "original_text": issue["suggested_anchor"],
                "proposed_replacement": "Table 1 presents patient outcomes across the study period.",
                "track_operation": issue["intro_operation"]}
        output, status = build_tracked_changes_document(word_bytes(doc), {"items": [item]})
        self.assertEqual(status["applied"], ["intro"])
        xml = Document(io.BytesIO(output))._element.body.xml
        self.assertLess(xml.index("Table 1 presents"), xml.index("Table 1. Patient outcomes"))
        self.assertIn("w:ins", xml)
        item["proposed_replacement"] = "Table 1 presents [describe the content accurately]."
        _, rejected = build_tracked_changes_document(word_bytes(doc), {"items": [item]})
        self.assertEqual(rejected["skipped"], ["intro"])

    def test_source_approval_is_not_recorded_as_applied_when_nothing_changes(self):
        doc = Document(); doc.add_paragraph("Smith (2023) informed the design.")
        doc.add_heading("References", level=1)
        doc.add_paragraph("Smith, J. (2023). Verified work. Journal of Care.")
        item = {"id": "source", "decision": "accepted", "category": "missing_reference",
                "original_text": "Smith (2023)", "proposed_replacement": "Smith (2023)",
                "secondary_replacement": "Smith, J. (2023). Verified work. Journal of Care.",
                "track_operation": "replace_citation_and_append_reference"}
        _, status = build_tracked_changes_document(word_bytes(doc), {"items": [item]})
        self.assertEqual(status["applied_count"], 0)
        self.assertIn("already present", status["unapplied"][0]["reason"])

    def test_unlocatable_approved_action_is_reported(self):
        doc = Document(); doc.add_paragraph("A paragraph that has no cited claim.")
        item = {"id": "source", "decision": "accepted", "category": "citation_needed",
                "original_text": "Missing passage", "proposed_replacement": "(Smith, 2023)",
                "track_operation": "insert_after"}
        _, status = build_tracked_changes_document(word_bytes(doc), {"items": [item]})
        self.assertEqual(status["applied_count"], 0)
        self.assertEqual(status["unapplied"][0]["id"], "source")

    def test_package_reports_track_changes_without_inserting_editorial_note(self):
        doc = Document(); doc.add_paragraph("A passage to revise.")
        item = {"id": "claim", "decision": "accepted", "category": "claim_support",
                "original_text": "A passage to revise.", "proposed_replacement": "A qualified passage.",
                "track_operation": "replace"}
        output, status = build_tracked_changes_document(word_bytes(doc), {"items": [item]})
        package = build_report_package("abc123", {"correction_plan": {"items": [item]},
                                                 "track_changes_application": status},
                                       {"CiteIntegrity_Track_Changes.docx": output})
        with zipfile.ZipFile(io.BytesIO(package)) as archive:
            report = archive.read("submission_readiness_report.html").decode()
            self.assertIn("change-control note", report)
            self.assertIn("1 of 1 accepted actions applied", report)
            self.assertIn("Applied", archive.read("correction_plan.csv").decode())
            word_xml = Document(io.BytesIO(archive.read("CiteIntegrity_Track_Changes.docx")))._element.body.xml
            self.assertNotIn("change-control note", word_xml)


if __name__ == "__main__":
    unittest.main()
