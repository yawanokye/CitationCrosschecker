"""Exercise the production package route with manuscripts lacking Word styles."""
import asyncio
import io
import json
import zipfile
from datetime import datetime
from unittest.mock import Mock
from urllib.parse import urlsplit

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import JSONResponse, Response
from starlette.background import BackgroundTask

from document_correction_pack import build_annotated_document, build_tracked_changes_document
from privacy_lifecycle import build_report_package
from test_approval_error_release import CLAIM, approval_payload, damage_checksum, document_bytes, result_fixture
from test_harvard_result_release import _load_result_route
from test_reference_review_release import document_xml, load_decision_route
from test_verification_coverage_release import load_functions


def manuscript(missing=True, wrong_type=False, bad_media=False):
    document = Document(io.BytesIO(document_bytes()))
    if missing or wrong_type:
        document.styles["Table Grid"].delete()
    if wrong_type:
        document.styles.add_style("Table Grid", WD_STYLE_TYPE.PARAGRAPH)
    stream = io.BytesIO()
    document.save(stream)
    raw = stream.getvalue()
    return damage_checksum(raw, "word/media/image1.png") if bad_media else raw


def load_routes(raw, result=None, delete_enabled=True):
    result = result or result_fixture()
    # Reuse the real result-preparation helper, with access/storage simulated.
    result_route, _ = _load_result_route(result=result)
    result_route.__globals__["get_access_mode"] = lambda *args: {"mode": "open_access"}
    purge = Mock(return_value={"ok": True})
    access = Mock()
    namespace = {
        "Request": Request, "HTTPException": HTTPException, "JSONResponse": JSONResponse,
        "Response": Response, "BackgroundTask": BackgroundTask,
        "now": lambda: datetime(2026, 10, 6).isoformat(),
        "_require_commercial_access": access,
        "load_job_record_fresh": lambda job: {"result": result, "file_name": "manuscript.docx"},
        "_prepare_student_result": result_route.__globals__["_prepare_student_result"],
        "redis_conn": Mock(get=Mock(return_value=raw)),
        "build_annotated_document": build_annotated_document,
        "build_tracked_changes_document": build_tracked_changes_document,
        "build_report_package": build_report_package,
        "DELETE_AFTER_PACKAGE_DOWNLOAD": delete_enabled,
        "_purge_job_content": purge,
    }
    load_functions("main.py", {"download_complete_report_package", "correction_application_status", "http_exception_handler"}, namespace)
    app = FastAPI()
    app.get("/api/report-package/{job_id}")(namespace["download_complete_report_package"])
    app.get("/api/corrections/{job_id}/application-status")(namespace["correction_application_status"])
    app.add_exception_handler(HTTPException, namespace["http_exception_handler"])
    return app, namespace, purge, access


def asgi_get(app, url, purge=None):
    parsed = urlsplit(url)
    messages = []
    async def receive():
        return {"type": "http.request", "body": b"", "more_body": False}
    async def send(message):
        if message["type"] == "http.response.body" and purge:
            assert not purge.called, "Content deletion must wait until response delivery"
        messages.append(message)
    scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": "GET",
             "scheme": "https", "path": parsed.path, "raw_path": parsed.path.encode(),
             "query_string": parsed.query.encode(), "headers": [],
             "server": ("test", 443), "client": ("test", 1234), "root_path": ""}
    asyncio.run(app(scope, receive, send))
    start = next(message for message in messages if message["type"] == "http.response.start")
    body = b"".join(message.get("body", b"") for message in messages if message["type"] == "http.response.body")
    return start["status"], {key.decode(): value.decode() for key, value in start["headers"]}, body


@pytest.mark.parametrize("missing,wrong_type", [(True, False), (False, True), (False, False)])
def test_annotated_report_handles_missing_or_redefined_table_style(missing, wrong_type):
    raw = manuscript(missing, wrong_type)
    result = result_fixture()
    from correction_plan import build_correction_plan
    output = build_annotated_document(raw, build_correction_plan(result))
    document = Document(io.BytesIO(output))
    assert document.paragraphs[0].text == CLAIM and len(document.tables) == 1
    assert len(document.tables[0].rows) > 1
    if missing or wrong_type:
        assert "<w:tblBorders>" in document_xml(output)
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(io.BytesIO(output)) as annotated:
        assert source.read("word/styles.xml") == annotated.read("word/styles.xml")
        assert source.read("word/media/image1.png") == annotated.read("word/media/image1.png")


def test_preflight_and_package_include_approved_tracked_edits_without_table_style():
    raw = manuscript(bad_media=True)
    result = result_fixture()
    decision, saved = load_decision_route(result, raw)
    async def payload():
        return approval_payload()
    asyncio.run(decision("job", Mock(json=payload)))
    assert saved.called
    app, _, purge, access = load_routes(raw, result)
    status, _, preflight = asgi_get(app, "/api/corrections/job/application-status")
    assert status == 200 and json.loads(preflight)["applied_count"] == 1
    status, headers, package = asgi_get(app, "/api/report-package/job?delete_after=0", purge)
    assert status == 200 and not purge.called and access.call_count == 2
    assert headers["content-type"] == "application/zip"
    assert headers["content-disposition"].startswith("attachment;")
    assert headers["x-content-deletion"] == "scheduled-expiry"
    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        assert archive.testzip() is None
        assert len(archive.namelist()) == 4
        tracked = archive.read("CiteIntegrity_Track_Changes.docx")
        annotated = archive.read("CiteIntegrity_Annotated_Manuscript.docx")
        assert document_xml(tracked).count("<w:ins ") == 2
        assert "(Smith, 2022)" in document_xml(tracked)
        assert len(Document(io.BytesIO(annotated)).tables) == 1
        assert "1 of 1 accepted actions applied" in archive.read("submission_readiness_report.html").decode()


@pytest.mark.parametrize("query,enabled,deleted", [("", True, False), ("?delete_after=0", True, False),
                                                  ("?delete_after=1", True, True), ("?delete_after=1", False, False)])
def test_deletion_follows_explicit_choice_and_successful_response(query, enabled, deleted):
    app, _, purge, _ = load_routes(manuscript(), delete_enabled=enabled)
    status, headers, _ = asgi_get(app, "/api/report-package/job" + query, purge)
    assert status == 200 and purge.called == deleted
    assert headers["x-content-deletion"] == ("after-download" if deleted else "scheduled-expiry")


@pytest.mark.parametrize("function,stage", [("build_annotated_document", "annotated manuscript"),
                                          ("build_tracked_changes_document", "tracked manuscript"),
                                          ("build_report_package", "report package")])
def test_package_failure_returns_actionable_error_and_does_not_delete(function, stage):
    app, namespace, purge, _ = load_routes(manuscript())
    namespace[function] = Mock(side_effect=RuntimeError("simulated failure"))
    status, headers, body = asgi_get(app, "/api/report-package/job?delete_after=1", purge)
    failure = json.loads(body)
    assert status == 500 and not purge.called
    assert failure["detail"] == failure["error"] and stage in failure["detail"]
    assert "has not been deleted" in failure["detail"]
    assert "simulated failure" not in failure["detail"]
    assert "content-disposition" not in headers


def test_deleted_analysis_cannot_be_downloaded_again():
    app, _, purge, _ = load_routes(manuscript(), {"result_deleted": True})
    status, _, body = asgi_get(app, "/api/report-package/job", purge)
    assert status == 410 and "deleted" in json.loads(body)["detail"] and not purge.called
