"""Generate reviewable DOCX outputs without silently changing scholarly content."""

from __future__ import annotations

import io
import re
import zipfile
from copy import deepcopy
from datetime import datetime, timezone
from typing import Any, Dict, List, Tuple

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def _docx_bytes(document: Document) -> bytes:
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _repack_docx_media_checksums(original_bytes: bytes) -> bytes:
    """Recover a media checksum error without changing any stored part bytes.

    Document XML, relationships and other parts must pass their normal CRC
    checks. A media member must still decompress in full. This does not repair
    missing/truncated ZIP members or alter images or manuscript text.
    """
    output = io.BytesIO()
    recovered = False
    with zipfile.ZipFile(io.BytesIO(original_bytes)) as source, zipfile.ZipFile(output, "w") as target:
        for info in source.infolist():
            try:
                data = source.read(info)
            except zipfile.BadZipFile as error:
                if not info.filename.startswith("word/media/") or not str(error).startswith("Bad CRC-32"):
                    raise
                with source.open(info) as member:
                    # CPython's ZIP stream supports suppressing only its CRC
                    # comparison. Decompression/truncation errors still raise.
                    member._expected_crc = None
                    data = member.read()
                if len(data) != info.file_size:
                    raise zipfile.BadZipFile("Incomplete media member")
                recovered = True
            target.writestr(deepcopy(info), data)
    if not recovered:
        raise zipfile.BadZipFile("No recoverable media checksum error")
    return output.getvalue()


def _load_docx_status(original_bytes: bytes | None):
    status = {"loaded": False, "media_checksum_repacked": False,
              "reason": "The original DOCX is no longer available. Upload the manuscript again to prepare tracked changes."}
    if original_bytes and original_bytes[:2] == b"PK":
        try:
            document = Document(io.BytesIO(original_bytes))
            return document, True, {**status, "loaded": True, "reason": ""}
        except zipfile.BadZipFile as error:
            if str(error).startswith("Bad CRC-32"):
                try:
                    recovered = _repack_docx_media_checksums(original_bytes)
                    document = Document(io.BytesIO(recovered))
                    return document, True, {"loaded": True, "media_checksum_repacked": True, "reason": ""}
                except Exception:
                    pass
            status["reason"] = "The original DOCX contains damaged or incomplete parts. Open it in Word, save a fresh DOCX copy and upload that copy before approving tracked changes."
        except Exception:
            status["reason"] = "The original DOCX could not be opened for tracked changes. Open it in Word, save a fresh DOCX copy and upload that copy."
    elif original_bytes:
        status["reason"] = "This upload is not an editable DOCX. Upload a Word DOCX version to apply document-level tracked changes."
    return Document(), False, status


def _load_docx(original_bytes: bytes | None) -> Tuple[Document, bool]:
    document, loaded, _ = _load_docx_status(original_bytes)
    return document, loaded


def _insert_annotation_after(paragraph, text: str) -> None:
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "B45309")
    bold = OxmlElement("w:b")
    rpr.extend([color, bold])
    t = OxmlElement("w:t")
    t.text = text
    r.extend([rpr, t])
    p.append(r)
    paragraph._p.addnext(p)


def build_annotated_document(original_bytes: bytes | None, plan: Dict[str, Any], original_name: str = "manuscript.docx") -> bytes:
    document, copied_original = _load_docx(original_bytes)
    if not copied_original:
        document.add_heading("CiteIntegrity Annotated Manuscript Review", level=1)
        document.add_paragraph(f"The original file {original_name} could not be reconstructed as DOCX. This document contains the complete annotation register.")
    items = plan.get("items") or []
    unmatched: List[Dict[str, Any]] = []
    for index, item in enumerate(items, 1):
        evidence = str(item.get("evidence") or "").strip()
        matched = False
        if copied_original and evidence:
            key = evidence[:120].lower()
            for paragraph in document.paragraphs:
                if key and key in paragraph.text.lower():
                    for run in paragraph.runs:
                        if key[:40] in run.text.lower():
                            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    _insert_annotation_after(
                        paragraph,
                        f"CiteIntegrity {item.get('id')}: {item.get('priority','').upper()} | {item.get('what_is_wrong') or item.get('title')} | Recommended action: {item.get('recommended_action')} | Confidence: {item.get('confidence','review required')} | Decision: {item.get('decision','pending')}",
                    )
                    matched = True
                    break
        if not matched:
            unmatched.append(item)

    document.add_page_break()
    document.add_heading("CiteIntegrity Correction Register", level=1)
    document.add_paragraph(plan.get("headline") or "Human review is required before submission.")
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["ID", "Priority", "Location", "What is wrong", "Why it matters", "Recommended correction", "Decision"]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for item in items:
        cells = table.add_row().cells
        values = [
            item.get("id"), item.get("priority"), str(item.get("location") or ""),
            item.get("what_is_wrong") or item.get("title"), item.get("why_it_matters"),
            item.get("recommended_action"), item.get("decision", "pending"),
        ]
        for cell, value in zip(cells, values):
            cell.text = str(value or "")
    document.add_paragraph("Annotations are advisory. New citations, source replacements, claims and interpretations require the student's explicit approval and verification.")
    return _docx_bytes(document)


def _normalised_raw_span(raw_text: str, target_text: str):
    """Map a whitespace-normalised target back to its exact Word text span."""
    target = " ".join(str(target_text or "").split())
    if not raw_text or not target:
        return None
    normalised_chars = []
    raw_positions = []
    in_space = False
    for raw_index, char in enumerate(raw_text):
        if char.isspace():
            if not in_space:
                normalised_chars.append(" ")
                raw_positions.append(raw_index)
            in_space = True
        else:
            normalised_chars.append(char)
            raw_positions.append(raw_index)
            in_space = False
    normalised = "".join(normalised_chars)
    start = normalised.find(target)
    if start < 0:
        return None
    end_index = start + len(target) - 1
    if end_index >= len(raw_positions):
        return None
    return raw_positions[start], raw_positions[end_index] + 1


def _tracked_replace(paragraph, original: str, replacement: str, change_id: int, replacement_segments=None) -> bool:
    full_text = paragraph.text
    start = full_text.find(original)
    if start < 0:
        span = _normalised_raw_span(full_text, original)
        if not span:
            return False
        start, end = span
    else:
        end = start + len(original)
    p = paragraph._p
    runs = [child for child in p if child.tag == qn("w:r")]
    # A hyperlink, field, drawing, break, or prior complex revision requires
    # more precise structural handling than a safe text replacement can offer.
    if any(child.tag not in {qn("w:rPr"), qn("w:t")} for run in runs for child in run):
        return False
    raw = "".join("".join(node.text or "" for node in run if node.tag == qn("w:t")) for run in runs)
    if raw != full_text:
        return False
    segments = replacement_segments if isinstance(replacement_segments, list) and replacement_segments else [{"text": replacement}]

    def copy_text_run(source, text):
        run = OxmlElement("w:r")
        props = source.find(qn("w:rPr"))
        if props is not None:
            run.append(deepcopy(props))
        node = OxmlElement("w:t"); node.set(qn("xml:space"), "preserve"); node.text = text
        run.append(node)
        return run

    stamp = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    deletion = OxmlElement("w:del"); deletion.set(qn("w:id"), str(change_id)); deletion.set(qn("w:author"), "CiteIntegrity"); deletion.set(qn("w:date"), stamp)
    insertion = OxmlElement("w:ins"); insertion.set(qn("w:id"), str(change_id + 1)); insertion.set(qn("w:author"), "CiteIntegrity"); insertion.set(qn("w:date"), stamp)
    offset = 0
    matches = []
    for run in runs:
        text = "".join(node.text or "" for node in run if node.tag == qn("w:t"))
        run_start, run_end = offset, offset + len(text)
        offset = run_end
        if run_start < end and run_end > start:
            matches.append((run, text, run_start, run_end))
    if not matches:
        return False
    for run, text, run_start, run_end in matches:
        left = max(0, start - run_start)
        right = min(len(text), end - run_start)
        if left:
            run.addprevious(copy_text_run(run, text[:left]))
        old_run = copy_text_run(run, text[left:right])
        old_run.find(qn("w:t")).tag = qn("w:delText")
        deletion.append(old_run)
        if run is matches[-1][0]:
            run.addprevious(deletion)
            # Word retains formatting around the changed span unless the
            # replacement explicitly requests italic reference segments.
            old_properties = matches[0][0].find(qn("w:rPr"))
            for segment in segments:
                text_value = str((segment or {}).get("text") or "")
                if not text_value:
                    continue
                new_run = OxmlElement("w:r")
                if old_properties is not None:
                    new_run.append(deepcopy(old_properties))
                if (segment or {}).get("italic"):
                    props = new_run.find(qn("w:rPr"))
                    if props is None:
                        props = OxmlElement("w:rPr"); new_run.insert(0, props)
                    props.append(OxmlElement("w:i"))
                    props.append(OxmlElement("w:iCs"))
                node = OxmlElement("w:t"); node.set(qn("xml:space"), "preserve"); node.text = text_value
                new_run.append(node); insertion.append(new_run)
            run.addprevious(insertion)
        if right < len(text):
            run.addprevious(copy_text_run(run, text[right:]))
        p.remove(run)
    return True


def _tracked_insert_after(paragraph, anchor: str, insertion_text: str, change_id: int) -> bool:
    full_text = paragraph.text
    start = full_text.find(anchor)
    if start < 0:
        span = _normalised_raw_span(full_text, anchor)
        if not span:
            return False
        start, split_at = span
    else:
        split_at = start + len(anchor)
    citation_insertion = insertion_text.strip().startswith(("(", "["))
    if citation_insertion and split_at and full_text[split_at - 1] in ".!?":
        # Place an author-year citation before the sentence's final punctuation.
        split_at -= 1
    p = paragraph._p
    runs = [child for child in p if child.tag == qn("w:r")]
    if any(child.tag not in {qn("w:rPr"), qn("w:t")} for run in runs for child in run):
        return False
    if "".join("".join(node.text or "" for node in run if node.tag == qn("w:t")) for run in runs) != full_text:
        return False
    stamp = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    insertion = OxmlElement("w:ins"); insertion.set(qn("w:id"), str(change_id)); insertion.set(qn("w:author"), "CiteIntegrity"); insertion.set(qn("w:date"), stamp)
    ir = OxmlElement("w:r"); it = OxmlElement("w:t"); it.set(qn("xml:space"), "preserve"); it.text = " " + insertion_text.strip(); ir.append(it); insertion.append(ir)
    offset = 0
    for run in runs:
        text = "".join(node.text or "" for node in run if node.tag == qn("w:t"))
        next_offset = offset + len(text)
        if split_at <= next_offset:
            within = split_at - offset
            if within == 0:
                run.addprevious(insertion)
            elif within == len(text):
                run.addnext(insertion)
            else:
                suffix = deepcopy(run)
                for node in suffix.findall(qn("w:t")):
                    suffix.remove(node)
                for node in run.findall(qn("w:t")):
                    run.remove(node)
                before = OxmlElement("w:t"); before.set(qn("xml:space"), "preserve"); before.text = text[:within]
                after = OxmlElement("w:t"); after.set(qn("xml:space"), "preserve"); after.text = text[within:]
                run.append(before); suffix.append(after)
                run.addnext(suffix)
                suffix.addprevious(insertion)
            return True
        offset = next_offset
    return False


def _editable_paragraphs(document: Document):
    """Include text in Word tables while keeping body paragraphs first."""
    yield from document.paragraphs
    seen = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if id(paragraph._p) not in seen:
                        seen.add(id(paragraph._p))
                        yield paragraph


def _tracked_paragraph_near(paragraph, text: str, change_id: int, *, before: bool = False) -> bool:
    if not text.strip() or re.search(r"\[(?:describe|insert|add|specify|verify)[^\]]*\]", text, re.I):
        return False  # Do not export an unfilled editorial placeholder.
    new_paragraph = OxmlElement("w:p")
    stamp = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    insertion = OxmlElement("w:ins")
    insertion.set(qn("w:id"), str(change_id)); insertion.set(qn("w:author"), "CiteIntegrity"); insertion.set(qn("w:date"), stamp)
    run = OxmlElement("w:r"); node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve"); node.text = text.strip()
    run.append(node); insertion.append(run); new_paragraph.append(insertion)
    if before:
        paragraph._p.addprevious(new_paragraph)
    else:
        paragraph._p.addnext(new_paragraph)
    return True


def _reference_section_anchor(document: Document):
    paragraphs = document.paragraphs
    heading_index = next((index for index, p in enumerate(paragraphs)
                          if p.text.strip().casefold() in {"references", "reference list", "bibliography"}), None)
    if heading_index is None:
        return None
    last = paragraphs[heading_index]
    for paragraph in paragraphs[heading_index + 1:]:
        style = paragraph.style.name if paragraph.style else ""
        if style.lower().startswith("heading") or paragraph.text.strip().casefold().startswith(("appendix", "supplementary materials")):
            break
        last = paragraph
    return last


def _tracked_append_reference(document: Document, reference_text: str, change_id: int) -> bool:
    if not reference_text.strip():
        return False
    anchor = _reference_section_anchor(document)
    return bool(anchor and _tracked_paragraph_near(anchor, reference_text, change_id))


def _original_revision_text(paragraph):
    """Recover the original anchor after another approved edit in this paragraph."""
    parts = []
    for child in paragraph._p:
        if child.tag == qn("w:ins"):
            continue
        tag = qn("w:delText") if child.tag == qn("w:del") else qn("w:t")
        parts.extend(node.text or "" for node in child.iter(tag))
    return "".join(parts)


def _tracked_reference_group(document, item, change_id):
    """Apply a reference and its approved citation edits together or not at all."""
    # Reload the current package so Document and its saved XML part share the
    # same tree. Deep-copying their proxies can discard earlier tracked edits.
    trial = Document(io.BytesIO(_docx_bytes(document)))
    original, replacement = str(item.get("original_text") or item.get("evidence") or ""), str(item.get("proposed_replacement") or "")
    targets = [p for p in _editable_paragraphs(trial) if original in p.text or _normalised_raw_span(p.text, original)]
    if len(targets) != 1 or not original or not replacement:
        return None, change_id, "The original reference could not be located uniquely."
    reference_target = targets[0]
    citation_targets = []
    for edit in item.get("related_edits") or []:
        anchor, old, new = str(edit.get("anchor") or ""), str(edit.get("original_text") or ""), str(edit.get("proposed_replacement") or "")
        matches = [p for p in _editable_paragraphs(trial) if p.text == anchor or " ".join(_original_revision_text(p).split()) == " ".join(anchor.split())]
        from reference_safety import is_review_instruction
        if len(matches) != 1 or matches[0]._p is reference_target._p or not old or not new or is_review_instruction(new):
            return None, change_id, "An associated in-text citation could not be located uniquely. No part of this approval was applied."
        expected = int(edit.get("expected_occurrences") or 1)
        if expected < 1 or " ".join(matches[0].text.split()).count(" ".join(old.split())) != expected:
            return None, change_id, "An associated citation changed since approval. No part of this approval was applied."
        citation_targets.append((matches[0], old, new, expected))
    changed = False
    if original != replacement:
        if not _tracked_replace(reference_target, original, replacement, change_id):
            return None, change_id, "The reference uses Word fields or complex structure that could not be edited safely."
        changed = True
        change_id += 2
    for paragraph, old, new, expected in citation_targets:
        for _ in range(expected):
            if not _tracked_replace(paragraph, old, new, change_id):
                return None, change_id, "An associated citation could not be tracked safely. No part of this approval was applied."
            changed = True
            change_id += 2
    return (trial if changed else None), change_id, "The approved reference and citations are already unchanged." if not changed else ""


def build_tracked_changes_document(original_bytes: bytes | None, plan: Dict[str, Any], original_name: str = "manuscript.docx") -> Tuple[bytes, Dict[str, Any]]:
    document, copied_original, document_status = _load_docx_status(original_bytes)
    applied: List[str] = []
    skipped: List[str] = []
    unapplied: List[Dict[str, str]] = []
    accepted = [item for item in (plan.get("items") or []) if item.get("decision") == "accepted"]
    if not copied_original:
        document.add_heading("CiteIntegrity Track Changes Report", level=1)
        document.add_paragraph(f"Track Changes could not be applied because {original_name} was not available as a valid DOCX file. Upload the DOCX version for document-level changes.")
        skipped = [str(item.get("id") or "") for item in accepted]
        unapplied = [{"id": item_id, "reason": document_status["reason"]} for item_id in skipped]
    else:
        def candidates(original):
            return [paragraph for paragraph in _editable_paragraphs(document)
                    if original in paragraph.text or _normalised_raw_span(paragraph.text, original)]

        def reference_present(text):
            expected = " ".join(text.split()).casefold()
            return any(" ".join("".join(p._p.itertext()).split()).casefold().find(expected) >= 0
                       for p in document.paragraphs)

        change_id = 1
        for item in accepted:
            original = str(item.get("original_text") or item.get("evidence") or "")
            replacement = str(item.get("proposed_replacement") or "")
            operation = str(item.get("track_operation") or "replace")
            from reference_safety import is_review_instruction
            if is_review_instruction(replacement) or is_review_instruction(item.get("secondary_replacement")):
                skipped.append(item.get("id")); unapplied.append({"id": item.get("id"), "reason": "A review instruction cannot be manuscript text."})
                continue
            changed = False
            reason = "The exact passage was not found uniquely in editable Word text. Select a longer anchor or review complex fields and tables."
            if operation == "replace_reference_and_citations":
                trial, next_id, reason = _tracked_reference_group(document, item, change_id)
                if trial is not None:
                    document, change_id, changed = trial, next_id, True
            elif operation in {"insert_after_and_append_reference", "replace_citation_and_append_reference"} and original and replacement:
                secondary = str(item.get("secondary_replacement") or "")
                targets = candidates(original)
                if len(targets) == 1 and (not secondary or reference_present(secondary) or _reference_section_anchor(document)):
                    if operation == "replace_citation_and_append_reference" and original == replacement:
                        primary_changed = bool(secondary and not reference_present(secondary))
                        if not primary_changed:
                            reason = "The citation and completed reference are already present; there is no tracked change to apply."
                    elif operation == "replace_citation_and_append_reference":
                        primary_changed = _tracked_replace(targets[0], original, replacement, change_id)
                        if primary_changed: change_id += 2
                    else:
                        primary_changed = _tracked_insert_after(targets[0], original, replacement, change_id)
                        if primary_changed: change_id += 1
                    if primary_changed:
                        if not secondary or reference_present(secondary):
                            changed = True
                        elif _tracked_append_reference(document, secondary, change_id):
                            changed = True; change_id += 1
                elif secondary and not _reference_section_anchor(document):
                    reason = "No References heading was found for the new entry. The citation was not inserted either."
            elif operation == "append_reference" and replacement:
                if not reference_present(replacement):
                    changed = _tracked_append_reference(document, replacement, change_id)
                else:
                    reason = "This completed reference is already present in the manuscript."
                if not changed and not _reference_section_anchor(document):
                    reason = "No References heading was found for a tracked reference insertion."
                if changed: change_id += 1
            elif operation in {"insert_paragraph_after", "insert_paragraph_before"} and original and replacement:
                targets = candidates(original)
                if len(targets) == 1:
                    changed = _tracked_paragraph_near(targets[0], replacement, change_id, before=operation == "insert_paragraph_before")
                if not changed and re.search(r"\[(?:describe|insert|add|specify|verify)[^\]]*\]", replacement, re.I):
                    reason = "Complete the proposed introduction before approving it."
                if changed: change_id += 1
            elif operation == "insert_after" and original and replacement:
                targets = candidates(original)
                if len(targets) == 1:
                    changed = _tracked_insert_after(targets[0], original, replacement, change_id)
                if changed: change_id += 1
            elif operation == "delete" and original:
                targets = candidates(original)
                if len(targets) == 1:
                    changed = _tracked_replace(targets[0], original, "", change_id)
                if changed: change_id += 2
            elif operation == "replace_all" and original and replacement:
                for paragraph in _editable_paragraphs(document):
                    # Each replacement removes the original from the ordinary
                    # runs; revision XML retains the deletion for Word review.
                    for _ in range(100):
                        if not _tracked_replace(paragraph, original, replacement, change_id):
                            break
                        changed = True; change_id += 2
            elif replacement:
                replacement_segments = ((item.get("supporting_metadata") or {}).get("reference_format_segments") if item.get("category") == "reference_style" else None)
                targets = candidates(original)
                if len(targets) == 1:
                    changed = _tracked_replace(targets[0], original, replacement, change_id, replacement_segments)
                if changed: change_id += 2
            if changed:
                applied.append(item.get("id"))
            else:
                skipped.append(item.get("id"))
                unapplied.append({"id": item.get("id"), "reason": reason})
    return _docx_bytes(document), {"applied": applied, "skipped": skipped, "unapplied": unapplied,
                                    "applied_count": len(applied), "accepted_count": len(accepted),
                                    "source_document": document_status}
